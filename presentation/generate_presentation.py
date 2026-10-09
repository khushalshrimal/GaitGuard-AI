"""
GaitGuard AI - PowerPoint Presentation Generator (Phase 18)
Generates GaitGuard_AI_Final_Presentation.pptx using python-pptx.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    EMERALD_DARK = RGBColor(6, 78, 59)      # #064e3b
    EMERALD_PRIMARY = RGBColor(16, 185, 129) # #10b981
    SLATE_DARK = RGBColor(15, 23, 42)       # #0f172a
    SLATE_MUTED = RGBColor(100, 116, 139)   # #64748b
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)

    blank_layout = prs.slide_layouts[6]

    slides_data = [
        {
            "title": "GaitGuard AI",
            "subtitle": "AI-Assisted Cattle Gait & Lameness-Risk Screening System\nVideo-Based Biomechanical Movement Analysis & Explainable AI Triage",
            "bg": EMERALD_DARK,
            "title_color": WHITE,
            "sub_color": RGBColor(167, 243, 208),
            "footer": "Phase 18 Final Release | Medical Disclaimer: Screening tool, not a veterinary diagnosis."
        },
        {
            "title": "Problem Statement & Impact",
            "subtitle": "Bovine Lameness: A Major Health & Economic Challenge in Cattle Farming",
            "points": [
                "Bovine lameness affects 25–35% of commercial dairy cattle worldwide.",
                "Causes severe animal pain, reduced milk yield, impaired fertility, and early culling.",
                "Traditional locomotion scoring is manual, subjective, and catches lameness at late stages.",
                "GaitGuard AI provides mobile-first automated screening to flag movement anomalies early."
            ]
        },
        {
            "title": "Proposed Solution & User Journey",
            "subtitle": "Simple 5-Step Mobile-First Screening Workflow",
            "points": [
                "1. Launch App: Open React PWA (Offline-ready with real-time API health status).",
                "2. Video Input: Drag & drop cattle walking video or record live via camera.",
                "3. Quality Gate: Automated check for motion blur, framing, and keypoint coverage.",
                "4. Calibrated Triage: Instant screening outcome (NORMAL / LAMENESS_RISK / INCONCLUSIVE).",
                "5. SHAP Evidence: Review biomechanical feature attributions and recommended next steps."
            ]
        },
        {
            "title": "Real Product Interface & User Experience",
            "subtitle": "Production-Hardened SaaS Interface Built for Field Usability",
            "points": [
                "Distinct visual cards for NORMAL, LAMENESS_RISK, and INCONCLUSIVE outcomes.",
                "Calibrated risk probability percentage with uncertainty bounds (±10%).",
                "Level 1 feature attributions and Level 2 domain gait metrics via SHAP.",
                "Zero PII collection and guaranteed zero temporary video persistence."
            ]
        },
        {
            "title": "System Architecture & Data Flow",
            "subtitle": "Modular Microservice Architecture (Client PWA -> FastAPI REST -> PyTorch Engine)",
            "points": [
                "Frontend PWA: React 18, TypeScript, Tailwind CSS, Vite.",
                "Backend REST API: FastAPI, Uvicorn, Pydantic v2, Request Correlation IDs (req_...).",
                "Security & Quality Gate: Magic-byte inspection (ftyp/RIFF), path traversal safety, OpenCV blur check.",
                "ML Engine: PyTorch BiLSTM neural network + Platt Sigmoid calibrator + SHAP Explainer."
            ]
        },
        {
            "title": "Keypoint Tracking & Gait Feature Engineering",
            "subtitle": "17 Anatomical Cattle Keypoints -> 76 Biomechanical Features",
            "points": [
                "Tracks 17 anatomical keypoints (withers, spine, hocks, fetlocks, hoofs).",
                "Savitzky-Golay polynomial trajectory filtering eliminates high-frequency tracking noise.",
                "Torso-length normalization makes gait metrics invariant to cow size and camera distance.",
                "76 features capture back-arch curvature, stance asymmetry, head nodding, and stride variance."
            ]
        },
        {
            "title": "BiLSTM Neural Network & Sequence Modeling",
            "subtitle": "Temporal Pattern Analysis Across 128-Frame Sequence Windows",
            "points": [
                "Bidirectional LSTM evaluates past and future temporal context in parallel.",
                "Sequence input tensor shape: (128, 76).",
                "Captures cyclic stride rhythm and dynamic motion traits impossible to detect in static 2D images."
            ]
        },
        {
            "title": "Internal Evaluation & Leak-Free GroupKFold",
            "subtitle": "Rigorous Animal-Level GroupKFold CV (0 Cow Overlap)",
            "points": [
                "Internal CV Accuracy: 81.97% | F1-Score: 0.7973 | Out-of-Fold ROC-AUC: 0.9016.",
                "Screening Recall (tau = 0.34): 89.15% screening sensitivity.",
                "GroupKFold Rigor: 0 animal overlap between training and validation folds in all 5 splits."
            ]
        },
        {
            "title": "Calibration, Triage & Explainable AI (SHAP)",
            "subtitle": "Platt Sigmoid Calibration & Deep SHAP Feature Attributions",
            "points": [
                "Platt Sigmoid Calibration maps raw neural net logits to calibrated screening probabilities.",
                "3-Way Triage Engine: Screening threshold tau = 0.34, inconclusive band [0.24, 0.44].",
                "Deep SHAP Explainer highlights feature contribution direction (+ Risk / - Neutralizer)."
            ]
        },
        {
            "title": "Privacy, Security & Data Protection",
            "subtitle": "Zero Temporary Video Persistence & Privacy-by-Design",
            "points": [
                "Temporary video files deleted immediately inside deterministic finally blocks after keypoint extraction.",
                "Magic-byte signature verification rejects renamed scripts or executables.",
                "Zero farmer PII collected; cryptographic request correlation tracking.",
                "Non-diagnostic disclaimer embedded across all API responses."
            ]
        },
        {
            "title": "Responsible Use & Field Validation Roadmap",
            "subtitle": "Responsible AI Deployment Policy & Preregistered Validation Protocol",
            "points": [
                "Internal CV metrics represent benchmark dataset performance, NOT clinical field sensitivity.",
                "External field validation remains BLOCKED — DATA COLLECTION REQUIRED.",
                "Preregistered 6-step protocol ready for independent blinded farm testing.",
                "Designed as a decision support tool to aid, not replace, veterinary professionals."
            ]
        },
        {
            "title": "Conclusion & Project Summary",
            "subtitle": "GaitGuard AI: Production-Ready Software System for Cattle Gait Screening",
            "points": [
                "Pipeline model latency: 66.59 ms (sub-100ms real-time target met).",
                "Automated test coverage: 239 backend tests & 13 frontend tests (100% pass rate).",
                "Production PWA build and non-root Docker deployment container ready.",
                "Thank you! Ready for Live Demo & Q&A."
            ]
        }
    ]

    for s_idx, s_data in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        
        # Background
        bg_color = s_data.get("bg", LIGHT_BG)
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color

        # Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
        tf = title_box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = s_data["title"]
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = s_data.get("title_color", SLATE_DARK)

        # Subtitle
        if "subtitle" in s_data:
            p2 = tf.add_paragraph()
            p2.text = s_data["subtitle"]
            p2.font.size = Pt(18)
            p2.font.color.rgb = s_data.get("sub_color", EMERALD_DARK if bg_color == LIGHT_BG else RGBColor(167, 243, 208))
            p2.space_before = Pt(6)

        # Points (Content)
        if "points" in s_data:
            content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5))
            ctf = content_box.text_frame
            ctf.word_wrap = True
            
            for p_idx, point_text in enumerate(s_data["points"]):
                cp = ctf.paragraphs[0] if p_idx == 0 else ctf.add_paragraph()
                cp.text = f"•  {point_text}"
                cp.font.size = Pt(20)
                cp.font.color.rgb = SLATE_DARK
                cp.space_before = Pt(14)

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.5))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.text = s_data.get("footer", "GaitGuard AI — Phase 18 Release | AI-Assisted Cattle Gait Screening")
        fp.font.size = Pt(12)
        fp.font.color.rgb = SLATE_MUTED if bg_color == LIGHT_BG else RGBColor(167, 243, 208)

    output_path = os.path.join("presentation", "GaitGuard_AI_Final_Presentation.pptx")
    os.makedirs("presentation", exist_ok=True)
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at {output_path}")

if __name__ == "__main__":
    create_presentation()
